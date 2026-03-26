from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

out = Path(r'C:\Users\USER\Documents\Python Project\movepal-ml-2\output\doc\MovePal_ML_Explainer.docx')
out.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('MovePal ML Prediction Service\nNon-Technical Explainer')
r.bold = True
r.font.size = Pt(16)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run('A plain-language explanation of what the system does, why it was built this way, and how it should be understood.').italic = True

def heading(text, level=1):
    doc.add_heading(text, level=level)

def para(text):
    doc.add_paragraph(text)

def bullets(items):
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

heading('1. What MovePal Does')
para('MovePal is a prediction service designed to estimate how crowded a Lagos BRT station is likely to be. Instead of only telling whether a road is busy, it tries to estimate whether a station is likely to be flowing, moderately busy, or heavily congested from a passenger point of view.')
para('The service returns one of three results: flowing, moderate, or heavy. These are operational states that help describe the expected crowd condition at a station.')
heading('2. Why This Problem Matters')
para('In Lagos, commuters often make travel decisions without knowing what station conditions will feel like when they arrive. A station may be easy to move through at one time of day and highly stressful at another. That uncertainty affects waiting time, comfort, route choice, and travel planning.')
para('The biggest challenge is that there is no clean public Lagos dataset that directly tells us how crowded each BRT station is at every point in time. That means the main machine learning problem is not only building a model, but also designing a reliable way to learn from limited and imperfect data.')
heading('3. The Core Idea Behind the Solution')
para('Instead of pretending that a perfect station-crowding dataset already exists, MovePal uses a practical approach called proxy learning. This means the system combines several real and structured signals that together can help estimate station congestion even when direct passenger counts are unavailable.')
para('In simple terms, the system asks: given what we know about the station itself, the time of day, the surrounding traffic pressure, and the weather, what is the most likely crowd state at this station right now?')
heading('4. The Three Data Layers')
heading('4.1 Station Master Layer', level=2)
para('This is the structural intelligence layer of the system. It contains information about the station itself, such as whether it is a terminal, whether it acts as a transfer point, how strategically important it is, and what type of area surrounds it.')
bullets([
    'station role: whether the station behaves like a terminal, interchange, hub, or standard stop',
    'busy factor: a prior intensity score showing how naturally busy the station is expected to be',
    'transfer score: an estimate of how much route-switching pressure exists at that station',
    'corridor type: whether the station sits on a primary, secondary, or local corridor',
    'land-use context: whether the area behaves more like a business, market, education, residential, or mixed-use zone',
])
heading('4.2 Live Observation Layer', level=2)
para('This layer adds real-time context. The system collects live traffic flow information and live weather information for each station area. These signals help explain the current environment around the station.')
bullets([
    'traffic pressure: how constrained movement is around the station compared with normal flow',
    'weather: rain, cloud cover, and related conditions that can influence waiting and movement behavior',
])
heading('4.3 Time Layer', level=2)
para('The system also uses time-based signals, because passenger movement changes across the day and across the week.')
bullets([
    'hour of day',
    'day of week',
    'rush-hour patterns',
    'late-night low-demand patterns',
])
heading('5. What Flowing, Moderate, and Heavy Mean')
para('These labels do not mean exact passenger counts. They describe operational crowd states.')
bullets([
    'Flowing: passengers can move through the station with little delay and minimal crowd buildup.',
    'Moderate: the station is visibly busy, but still manageable. Waiting and movement pressure are noticeable.',
    'Heavy: there is strong crowd buildup, longer queues, slower movement, and greater passenger pressure.',
])
heading('6. How the Model Learns')
para('The system collects observations for the stations and combines them with station-level structural features. From those observations, it builds a training dataset. That dataset is then used to train a machine learning model.')
para('The chosen model is a Random Forest classifier built with scikit-learn. It was selected because this is a structured tabular prediction problem, not an image or language problem. Random Forest works well when there are many interacting factors, such as station type, time, traffic pressure, and weather.')
heading('7. Why Scikit-Learn Was Used')
para('Scikit-learn was used because it is reliable, well tested, fast to build with, and appropriate for structured machine learning tasks like this one. It also provides strong evaluation tools, which made it easier to validate the model under hackathon time constraints.')
heading('8. Why This Approach Is Better Than the Old Fake-Data Version')
para('The older version mainly relied on fabricated labels and assumed rules. That kind of model can appear to perform well, but in reality it mostly learns the assumptions that were manually written into the data.')
para('The upgraded version is stronger because it combines real station locations, live external signals, and structured station intelligence. In other words, it does not claim to have perfect crowd truth, but it is much closer to reality than a purely synthetic approach.')
heading('9. What the Results Mean')
para('The latest training run used 22,908 model-ready rows and achieved 96.99% accuracy on the held-out test split. This is a strong result for a hackathon prototype, especially because the class balance was improved so that flowing, moderate, and heavy conditions are all represented more meaningfully.')
para('A learning curve was also generated using scikit-learn. The validation accuracy improved as more data was used, and the gap between training and validation became smaller. This suggests that the model is learning useful patterns rather than simply memorizing a lucky split.')
heading('10. What the System Is and Is Not')
para('It is important to describe the system honestly.')
bullets([
    'It is a proxy-based passenger congestion predictor.',
    'It is not a direct headcount system.',
    'It uses real and engineered signals to estimate crowd state.',
    'It is designed to improve as more observations are collected.',
])
heading('11. What Happens After the Hackathon')
para('The current system is already deployable, but it is also designed to grow.')
bullets([
    'More station snapshots can be collected over time.',
    'The model can be retrained as the dataset grows.',
    'Crowdsourced commuter reports can later be added as stronger ground-truth feedback.',
    'A production version can move collection and retraining into scheduled cloud workflows.',
])
heading('12. Simple Final Summary')
para('MovePal solves a difficult Lagos mobility problem by combining station knowledge, live contextual data, and machine learning. Because a perfect public station-crowding dataset does not exist, the project focuses on building a realistic proxy-learning pipeline. The result is a practical congestion prediction service that works today and becomes more useful as more real observations are collected.')

doc.save(out)
print(out)
