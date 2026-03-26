from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

base = Path(r'C:\Users\USER\Documents\Python Project\movepal-ml-2\output\doc')
base.mkdir(parents=True, exist_ok=True)

def setup_doc(title, subtitle):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(16)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.add_run(subtitle).italic = True
    return doc

judge = setup_doc('MovePal ML Technical Brief', 'A concise judge-facing summary of the problem, method, validation, and limitations.')
judge.add_heading('1. Problem Context', level=1)
judge.add_paragraph('Lagos lacks a clean public dataset of station-level passenger crowding for BRT operations. This makes direct supervised learning difficult. Most simple demo systems respond by fabricating labels, which creates weak credibility. MovePal treats the absence of ground-truth crowd labels as the real machine learning problem.')
judge.add_heading('2. Core ML Framing', level=1)
judge.add_paragraph('The system uses a proxy-learning pipeline. Instead of claiming direct passenger-count ground truth, it combines structured station priors, live traffic flow, weather context, and time features to generate operational congestion states: flowing, moderate, and heavy.')
judge.add_heading('3. Data Layers', level=1)
for item in [
    'Station master layer: station role, terminal/interchange status, busy factor, transfer score, corridor type, land-use context.',
    'Observation layer: live TomTom traffic flow and Open-Meteo weather.',
    'Training layer: model-ready rows built from observed and expanded time-context combinations.',
]:
    judge.add_paragraph(item, style='List Bullet')
judge.add_heading('4. Model Choice', level=1)
judge.add_paragraph('A Random Forest classifier in scikit-learn was used because this is a structured tabular classification task with nonlinear feature interactions. The model is fast to train, robust, explainable, and suitable for hackathon time constraints.')
judge.add_heading('5. Validation', level=1)
judge.add_paragraph('Latest training run: 22,908 rows, 96.99% held-out test accuracy. Learning-curve validation showed training accuracy remaining stable around 97% while validation accuracy improved from about 89.9% to 96.8% as training size increased. This supports that the model is learning stable structure rather than relying on a lucky split.')
judge.add_heading('6. Why This Is Stronger Than the Synthetic Version', level=1)
judge.add_paragraph('The previous version learned from fabricated rules. The upgraded version learns from validated station metadata and live contextual signals. The novelty is not only the classifier but the weak-supervision style dataset-generation pipeline itself.')
judge.add_heading('7. Limits and Honest Framing', level=1)
for item in [
    'This is a proxy-based passenger congestion predictor, not a direct headcount system.',
    'Some station priors are engineered rather than officially annotated.',
    'The system is designed to improve as more observation windows and eventual crowdsourced reports are added.',
]:
    judge.add_paragraph(item, style='List Bullet')
judge_path = base / 'MovePal_ML_Judge_Brief.docx'
judge.save(judge_path)

colleague = setup_doc('MovePal ML Presentation Notes', 'A short handoff version for teammates to quickly understand and present the system.')
colleague.add_heading('What It Is', level=1)
colleague.add_paragraph('MovePal predicts how crowded a Lagos BRT station is likely to be: flowing, moderate, or heavy.')
colleague.add_heading('Why It Matters', level=1)
colleague.add_paragraph('Commuters usually do not know station conditions before arrival. That creates uncertainty in waiting time, comfort, and route planning.')
colleague.add_heading('What Makes It Different', level=1)
colleague.add_paragraph('Instead of using a fake crowding dataset, the system uses a station intelligence layer plus live traffic and weather context to create a proxy congestion dataset.')
colleague.add_heading('How It Works', level=1)
for item in [
    'Use station master data to describe each BRT station structurally.',
    'Collect live traffic and weather signals.',
    'Combine them with time features.',
    'Train a Random Forest model to classify flowing, moderate, or heavy.',
    'Serve predictions through a Flask API.',
]:
    colleague.add_paragraph(item, style='List Bullet')
colleague.add_heading('Current Results', level=1)
colleague.add_paragraph('Latest run: 22,908 training rows, 96.99% test accuracy, and a learning curve showing validation improving as more data is added.')
colleague.add_heading('Safe Talking Points', level=1)
for item in [
    'This is a practical proxy-learning system for a city with limited ground-truth crowding data.',
    'It is deployable now and designed to improve with more observations.',
    'It should be described as a passenger-congestion estimator, not a perfect headcount system.',
]:
    colleague.add_paragraph(item, style='List Bullet')
colleague_path = base / 'MovePal_ML_Presentation_Notes.docx'
colleague.save(colleague_path)

print(judge_path)
print(colleague_path)
